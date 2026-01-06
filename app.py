# app.py  (v0.6.8)  v0.6.3 UI 유지 + 5개 회사 입력/합산 + LEGACY 제거 + DB 정리
# ✅ 변경: 기존 BEP(영업이익=0) 제거, "이자 포함 BEP"만 표시/사용

import streamlit as st
import pandas as pd
import sqlite3
import json
import io
import os
from dataclasses import dataclass
from datetime import datetime

# ReportLab (PDF)
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Optional PDF merge (if installed)
try:
    from PyPDF2 import PdfReader, PdfWriter
    HAS_PYPDF2 = True
except Exception:
    HAS_PYPDF2 = False


# =============================
# App Meta
# =============================
APP_VERSION = "0.6.8"  # ✅ 기존 BEP 제거, 이자 포함 BEP만
DEVELOPER_NAME = "한동석"
COPYRIGHT_TEXT = f"© {datetime.now().year} {DEVELOPER_NAME}. All rights reserved."
DB_PATH = "mvp_finance.db"

st.set_page_config(page_title="원가·손익·BEP 경영분석", layout="wide")


# =============================
# 회사 설정 (✅ LEGACY 제거)
# =============================
COMPANIES = [
    "신한포장",
    "진성디앤피",
    "화진포장",
    "성화티앤피",
    "도희팩",
]
GROUP_OPTION = "전체(합산)"


# =============================
# Session helpers
# =============================
def period_key(company: str, year: int, month: int) -> str:
    return f"{company}::{int(year)}::{int(month)}"


def init_work_state_for_period(pkey: str, defaults: dict):
    last = st.session_state.get("_work_period_key")
    if last == pkey:
        return

    st.session_state["sga_rows_work"] = (defaults.get("sga_rows") or build_default_sga_rows())
    st.session_state["nonop_income_work"] = (defaults.get("nonop_income_rows") or build_default_nonop_rows(DEFAULT_NONOP_INCOME))
    st.session_state["nonop_expense_work"] = (defaults.get("nonop_expense_rows") or build_default_nonop_rows(DEFAULT_NONOP_EXPENSE))

    st.session_state["_work_period_key"] = pkey


# =============================
# DB (회사+연도+월)
# =============================
def db_conn():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    return conn


def _table_has_company_column(conn: sqlite3.Connection) -> bool:
    try:
        cur = conn.execute("PRAGMA table_info(periods)")
        cols = [r[1] for r in cur.fetchall()]
        return "company" in cols
    except Exception:
        return False


def ensure_db_schema(conn: sqlite3.Connection):
    cur = conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='periods'")
    exists = cur.fetchone() is not None

    if exists and (not _table_has_company_column(conn)):
        conn.execute("DROP TABLE IF EXISTS periods")
        conn.commit()

    conn.execute("""
    CREATE TABLE IF NOT EXISTS periods (
        company TEXT NOT NULL,
        year INTEGER NOT NULL,
        month INTEGER NOT NULL,
        payload TEXT NOT NULL,
        updated_at TEXT NOT NULL,
        PRIMARY KEY (company, year, month)
    )
    """)
    conn.commit()

    placeholders = ",".join(["?"] * len(COMPANIES))
    conn.execute(f"DELETE FROM periods WHERE company NOT IN ({placeholders})", COMPANIES)
    conn.commit()


def purge_all_data():
    conn = db_conn()
    ensure_db_schema(conn)
    conn.execute("DELETE FROM periods")
    conn.commit()
    conn.close()


def save_period(company: str, year: int, month: int, payload: dict):
    if company not in COMPANIES:
        raise ValueError("회사명이 유효하지 않습니다.")
    conn = db_conn()
    ensure_db_schema(conn)
    conn.execute(
        "INSERT OR REPLACE INTO periods(company, year, month, payload, updated_at) VALUES (?,?,?,?,?)",
        (company, year, month, json.dumps(payload, ensure_ascii=False), datetime.now().isoformat(timespec="seconds"))
    )
    conn.commit()
    conn.close()


def load_period(company: str, year: int, month: int) -> dict:
    if company not in COMPANIES:
        return {}
    conn = db_conn()
    ensure_db_schema(conn)
    cur = conn.execute("SELECT payload FROM periods WHERE company=? AND year=? AND month=?", (company, year, month))
    row = cur.fetchone()
    conn.close()
    if not row:
        return {}
    try:
        return json.loads(row[0])
    except Exception:
        return {}


def list_periods(company: str | None = None):
    conn = db_conn()
    ensure_db_schema(conn)
    if company and company in COMPANIES:
        cur = conn.execute(
            "SELECT company, year, month, updated_at FROM periods WHERE company=? ORDER BY year DESC, month DESC",
            (company,)
        )
    else:
        cur = conn.execute("SELECT company, year, month, updated_at FROM periods ORDER BY year DESC, month DESC")
    rows = cur.fetchall()
    conn.close()
    return rows


def load_periods_range(company: str, year: int, months: list[int]) -> list[dict]:
    if company not in COMPANIES or not months:
        return []

    conn = db_conn()
    ensure_db_schema(conn)

    placeholders = ",".join(["?"] * len(months))
    query = f"""
        SELECT month, payload
        FROM periods
        WHERE company=? AND year=? AND month IN ({placeholders})
        ORDER BY month
    """
    cur = conn.execute(query, [company, year, *months])
    rows = cur.fetchall()
    conn.close()

    out = []
    for m, payload_json in rows:
        try:
            payload = json.loads(payload_json)
        except Exception:
            payload = {}
        payload["_year"] = int(year)
        payload["_month"] = int(m)
        out.append(payload)
    return out


def load_periods_range_group(year: int, months: list[int], companies: list[str]) -> list[dict]:
    if not months:
        return []
    companies = [c for c in companies if c in COMPANIES]
    if not companies:
        return []

    conn = db_conn()
    ensure_db_schema(conn)

    placeholders_m = ",".join(["?"] * len(months))
    placeholders_c = ",".join(["?"] * len(companies))
    query = f"""
        SELECT company, month, payload
        FROM periods
        WHERE year=? AND month IN ({placeholders_m}) AND company IN ({placeholders_c})
        ORDER BY month, company
    """
    cur = conn.execute(query, [year, *months, *companies])
    rows = cur.fetchall()
    conn.close()

    by_month = {int(m): [] for m in months}
    for c, m, payload_json in rows:
        try:
            payload = json.loads(payload_json)
        except Exception:
            payload = {}
        by_month[int(m)].append((str(c), payload))

    out = []
    for m in months:
        out.append({"_year": int(year), "_month": int(m), "_payloads": by_month.get(int(m), [])})
    return out


def delete_period(company: str, year: int, month: int):
    if company not in COMPANIES:
        return
    conn = db_conn()
    ensure_db_schema(conn)
    conn.execute("DELETE FROM periods WHERE company=? AND year=? AND month=?", (company, year, month))
    conn.commit()
    conn.close()


def delete_period_all_companies(year: int, month: int):
    conn = db_conn()
    ensure_db_schema(conn)
    conn.execute("DELETE FROM periods WHERE year=? AND month=?", (year, month))
    conn.commit()
    conn.close()


# =============================
# Tax (KR)
# =============================
def corp_tax_national_kr(tax_base_krw: float) -> float:
    tb = max(0.0, float(tax_base_krw))
    if tb <= 200_000_000:
        return tb * 0.09
    elif tb <= 20_000_000_000:
        return tb * 0.19 - 20_000_000
    elif tb <= 300_000_000_000:
        return tb * 0.21 - 420_000_000
    else:
        return tb * 0.24 - 9_420_000_000


def corp_tax_total_kr(tax_base_krw: float, include_local: bool = True) -> dict:
    national = corp_tax_national_kr(tax_base_krw)
    local = national * 0.10 if include_local else 0.0
    total = national + local
    return {
        "tax_base": max(0.0, float(tax_base_krw)),
        "national_cit": max(0.0, national),
        "local_income_tax": max(0.0, local),
        "total_tax": max(0.0, total),
        "assumption": "지방소득세 = 법인세(국세) × 10% (경영관리 추정)"
    }


def estimate_monthly_tax_from_pretax_annualized(monthly_pretax: float, include_local: bool = True) -> dict:
    pretax_m = float(monthly_pretax)
    if pretax_m <= 0:
        return {
            "annualized_tax_base": 0.0,
            "monthly_tax": 0.0,
            "detail": corp_tax_total_kr(0.0, include_local=include_local),
            "note": "당월 이익이 0 이하 → 월 법인세 0(추정)"
        }
    annual_base = pretax_m * 12.0
    detail = corp_tax_total_kr(annual_base, include_local=include_local)
    monthly_tax = detail["total_tax"] / 12.0
    return {
        "annualized_tax_base": annual_base,
        "monthly_tax": monthly_tax,
        "detail": detail,
        "note": "연환산(월×12) 기준 월추정 법인세"
    }


# =============================
# Helpers
# =============================
def safe_div(a: float, b: float) -> float:
    return a / b if b else 0.0


def grade_cm_ratio(cm_ratio: float) -> str:
    if cm_ratio >= 0.30:
        return "안정"
    if cm_ratio >= 0.20:
        return "주의"
    if cm_ratio >= 0.10:
        return "위험"
    return "심각"


def to_money(x: float) -> str:
    return f"{x:,.0f}"


def ratio_str(x: float) -> str:
    return f"{x * 100:,.1f}%"


def register_korean_font_if_possible() -> str:
    candidates = [
        r"C:\Windows\Fonts\malgun.ttf",
        r"C:\Windows\Fonts\Malgun.ttf",
        r"C:\Windows\Fonts\malgunsl.ttf",
        r"C:\Windows\Fonts\NanumGothic.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                font_name = "KOR_FONT"
                pdfmetrics.registerFont(TTFont(font_name, path))
                return font_name
            except Exception:
                pass
    return "Helvetica"


def get_dm_used_only(cost_items: dict) -> float:
    try:
        return float(cost_items.get("직접재료비", {}).get("당기사용원재료", 0.0) or 0.0)
    except Exception:
        return 0.0


def months_for_period(selected_month: int, kind: str) -> list[int]:
    m = int(selected_month)
    if kind == "월별(선택월)":
        return [m]
    if kind == "분기누적(QTD)":
        q = (m - 1) // 3 + 1
        start = (q - 1) * 3 + 1
        return list(range(start, m + 1))
    if kind == "연도누적(YTD)":
        return list(range(1, m + 1))
    return list(range(1, 13))


# =============================
# Default templates
# =============================
DEFAULT_COST_ITEMS = {
    "직접재료비": [
        "기초원재료",
        "당기원재료입고",
        "당기사용원재료",
        "기말원재료",
    ],
    "직접노무비": [
        "생산직 급여",
        "생산직 상여/수당",
        "생산직 4대보험(회사부담)",
        "생산직 퇴직급여충당금",
        "기타 직접노무"
    ],
    "제조간접비": [
        "부재료/잉크/접착제/철선 등(소모성)",
        "포장재/부포재/랩/밴드 등(소모성)",
        "외주가공비",
        "공정 소모품/부품/보수자재",
        "공무/현장관리 급여",
        "공무/현장관리 4대보험",
        "공무/현장관리 퇴직급여충당금",
        "전기료",
        "가스/연료비",
        "수도/환경비",
        "수선비/정비비",
        "공구기구비/소모품",
        "안전/방역/청소비",
        "공장 임차료",
        "운반구렌탈료",
        "보험료(공장/설비)",
        "차량관리비(화물차)",
        "운반비(공장 내/외)",
        "기계.설비 감가상각비",
        "공장 건물 감가상각비",
        "기타 제조간접비"
    ]
}

DEFAULT_SGA_ITEMS = [
    ("사무직 급여", "고정"),
    ("상여/수당/포상", "고정"),
    ("4대보험 회사부담", "고정"),
    ("퇴직급여충당금", "고정"),
    ("복리후생비", "고정"),
    ("사무용 소모품", "고정"),
    ("사무실 임차료/렌탈비", "고정"),
    ("통신비", "고정"),
    ("광고선전비", "변동"),
    ("판매수수료", "변동"),
    ("물류/배송비", "변동"),
    ("차량 렌트료(영업)", "고정"),
    ("차량 관리비(영업)", "고정"),
    ("지급수수료", "변동"),
    ("세금과공과", "고정"),
    ("대손상각비", "변동"),
    ("접대비", "고정"),
    ("출장비", "변동"),
    ("도서/자료 구입비", "변동"),
    ("특허/상표 유지비", "고정"),
    ("보험료(사무/영업)", "고정"),
    ("소프트웨어 사용료", "고정"),
    ("클라우드 서버 비용", "고정"),
    ("홈페이지 유지비", "고정"),
    ("사무실 비품 감가상각비", "고정"),
    ("사무실 건물 감가상각비", "고정"),
    ("영업용 차량 감가상각비", "고정"),
]

DEFAULT_NONOP_INCOME = ["이자수익", "환차익", "자산처분이익", "잡이익(파지매출 등)"]
DEFAULT_NONOP_EXPENSE = ["이자비용", "환차손", "자산처분손실", "잡손실"]


def normalize_sga_type(typ: str) -> str:
    return "변동" if typ == "변동" else "고정"


def build_default_cost_items():
    cost = {}
    for sec, items in DEFAULT_COST_ITEMS.items():
        cost[sec] = {it: 0.0 for it in items}
    return cost


def build_default_sga_rows():
    return [{"item": name, "amount": 0.0, "type": typ} for name, typ in DEFAULT_SGA_ITEMS]


def build_default_nonop_rows(items):
    return [{"item": it, "amount": 0.0} for it in items]


# =============================
# Core compute
# =============================
@dataclass
class PeriodInput:
    sales: float
    beg_fg: float
    end_fg: float
    cost_items: dict
    sga_rows: list
    nonop_income_rows: list
    nonop_expense_rows: list
    tax_mode: str
    tax_manual: float
    include_local_tax: bool


def sum_cost_section(cost_items: dict, section: str) -> float:
    sec = cost_items.get(section, {})
    return float(sum(float(v or 0) for v in sec.values()))


def compute_all(p: PeriodInput) -> dict:
    dm = get_dm_used_only(p.cost_items)
    dl = sum_cost_section(p.cost_items, "직접노무비")
    moh = sum_cost_section(p.cost_items, "제조간접비")
    cogm = dm + dl + moh

    cogs = p.beg_fg + cogm - p.end_fg

    sga_fixed = sum(float(r["amount"]) for r in p.sga_rows if r["type"] == "고정")
    sga_variable = sum(float(r["amount"]) for r in p.sga_rows if r["type"] == "변동")
    sga_total = sga_fixed + sga_variable

    nonop_income = sum(float(r["amount"]) for r in p.nonop_income_rows)
    nonop_expense = sum(float(r["amount"]) for r in p.nonop_expense_rows)

    gross_profit = p.sales - cogs
    op_profit = gross_profit - sga_total
    pretax = op_profit + nonop_income - nonop_expense

    tax_info = None
    if p.tax_mode == "AUTO":
        tax_info = estimate_monthly_tax_from_pretax_annualized(pretax, include_local=p.include_local_tax)
        tax = tax_info["monthly_tax"]
    elif p.tax_mode == "MANUAL":
        tax = p.tax_manual
    else:
        tax = 0.0

    net_income = pretax - tax

    # ✅ 공헌이익 / 이자 포함 BEP
    variable_total = cogs + sga_variable
    cm = p.sales - variable_total
    cm_ratio = safe_div(cm, p.sales)

    # ✅ 이자(영업외비용)까지 커버하는 BEP 매출
    interest_bep_sales = ((sga_fixed + nonop_expense) / cm_ratio) if cm_ratio > 0 else 0.0

    return {
        "dm": dm, "dl": dl, "moh": moh,
        "cogm": cogm,
        "cogs": cogs,
        "gross_profit": gross_profit,
        "sga_fixed": sga_fixed,
        "sga_variable": sga_variable,
        "sga_total": sga_total,
        "op_profit": op_profit,
        "nonop_income": nonop_income,
        "nonop_expense": nonop_expense,
        "pretax": pretax,
        "tax": tax,
        "tax_info": tax_info,
        "net_income": net_income,
        "variable_total": variable_total,
        "cm": cm,
        "cm_ratio": cm_ratio,
        "grade": grade_cm_ratio(cm_ratio),

        # ✅ 이자 포함 BEP만 제공
        "interest_bep_sales": interest_bep_sales,
        "gap_sales_vs_interest_bep": p.sales - interest_bep_sales,

        "op_margin": safe_div(op_profit, p.sales),
        "gross_margin": safe_div(gross_profit, p.sales),
    }


def payload_to_period_input(payload: dict) -> PeriodInput:
    raw_cost = payload.get("cost_items", {}) or build_default_cost_items()
    raw_sga = payload.get("sga_rows", []) or build_default_sga_rows()

    sga = []
    for r in raw_sga:
        sga.append({
            "item": r.get("item", ""),
            "amount": float(r.get("amount", 0.0)),
            "type": normalize_sga_type(r.get("type", "고정"))
        })

    return PeriodInput(
        sales=float(payload.get("sales", 0.0)),
        beg_fg=float(payload.get("beg_fg", 0.0)),
        end_fg=float(payload.get("end_fg", 0.0)),
        cost_items=raw_cost,
        sga_rows=sga,
        nonop_income_rows=[{"item": r.get("item", ""), "amount": float(r.get("amount", 0.0))}
                           for r in payload.get("nonop_income_rows", build_default_nonop_rows(DEFAULT_NONOP_INCOME))],
        nonop_expense_rows=[{"item": r.get("item", ""), "amount": float(r.get("amount", 0.0))}
                            for r in payload.get("nonop_expense_rows", build_default_nonop_rows(DEFAULT_NONOP_EXPENSE))],
        tax_mode=str(payload.get("tax_mode", "AUTO")),
        tax_manual=float(payload.get("tax_manual", 0.0)),
        include_local_tax=bool(payload.get("include_local_tax", True)),
    )


# =============================
# ✅ 합산 로직
# =============================
def merge_payloads_for_group(payloads: list[tuple[str, dict]]) -> dict:
    merged = {
        "sales": 0.0,
        "beg_fg": 0.0,
        "end_fg": 0.0,
        "cost_items": build_default_cost_items(),
        "sga_rows": [],
        "nonop_income_rows": [],
        "nonop_expense_rows": [],
        "tax_mode": "OFF",
        "tax_manual": 0.0,
        "include_local_tax": True,
        "_tax_sum_override": 0.0,
        "_tax_info_note": "합산 세금 = 회사별 산출 tax 합계(AUTO/MANUAL/OFF 포함)",
        "_companies_included": [],
    }

    sga_map = {}
    noi_map = {}
    noe_map = {}

    tax_sum = 0.0
    included = []

    for company, payload in payloads:
        if company not in COMPANIES:
            continue
        if not payload:
            continue

        included.append(company)

        p = payload_to_period_input(payload)
        out = compute_all(p)

        merged["sales"] += p.sales
        merged["beg_fg"] += p.beg_fg
        merged["end_fg"] += p.end_fg

        for sec in ["직접재료비", "직접노무비", "제조간접비"]:
            sec_dict = payload.get("cost_items", {}).get(sec, {}) if payload.get("cost_items") else {}
            if not isinstance(sec_dict, dict):
                sec_dict = {}
            for k, v in sec_dict.items():
                merged["cost_items"][sec][k] = float(merged["cost_items"][sec].get(k, 0.0)) + float(v or 0.0)

        for r in payload.get("sga_rows", []) or []:
            item = str(r.get("item", "")).strip()
            typ = normalize_sga_type(str(r.get("type", "고정")))
            amt = float(r.get("amount", 0.0) or 0.0)
            if not item:
                continue
            sga_map[(item, typ)] = float(sga_map.get((item, typ), 0.0)) + amt

        for r in payload.get("nonop_income_rows", []) or []:
            item = str(r.get("item", "")).strip()
            amt = float(r.get("amount", 0.0) or 0.0)
            if item:
                noi_map[item] = float(noi_map.get(item, 0.0)) + amt

        for r in payload.get("nonop_expense_rows", []) or []:
            item = str(r.get("item", "")).strip()
            amt = float(r.get("amount", 0.0) or 0.0)
            if item:
                noe_map[item] = float(noe_map.get(item, 0.0)) + amt

        tax_sum += float(out.get("tax", 0.0) or 0.0)

    merged["sga_rows"] = [{"item": k[0], "type": k[1], "amount": float(v)} for k, v in sga_map.items()]
    merged["nonop_income_rows"] = [{"item": k, "amount": float(v)} for k, v in noi_map.items()]
    merged["nonop_expense_rows"] = [{"item": k, "amount": float(v)} for k, v in noe_map.items()]

    merged["_tax_sum_override"] = tax_sum
    merged["_companies_included"] = included
    return merged


def compute_all_group(payloads: list[tuple[str, dict]]) -> tuple[PeriodInput, dict]:
    merged_payload = merge_payloads_for_group(payloads)
    p = payload_to_period_input(merged_payload)
    out = compute_all(p)

    tax_sum = float(merged_payload.get("_tax_sum_override", 0.0) or 0.0)
    out["tax"] = tax_sum
    out["tax_info"] = {
        "annualized_tax_base": 0.0,
        "monthly_tax": tax_sum,
        "detail": {"national_cit": 0.0, "local_income_tax": 0.0, "total_tax": tax_sum,
                   "assumption": merged_payload.get("_tax_info_note", "")},
        "note": merged_payload.get("_tax_info_note", ""),
    }
    out["net_income"] = out["pretax"] - out["tax"]
    out["gap_sales_vs_interest_bep"] = p.sales - out["interest_bep_sales"]
    return p, out


# =============================
# DataFrames
# =============================
def df_cost_statement(cost_items: dict, cogm: float) -> pd.DataFrame:
    rows = []

    dm_sec = cost_items.get("직접재료비", {})

    beg = float(dm_sec.get("기초원재료", 0.0) or 0.0)
    inbound_default = dm_sec.get("당기원재료입고", None)
    if inbound_default is None:
        inbound_default = dm_sec.get("당기원재료매입", 0.0)
    inbound = float(inbound_default or 0.0)

    used = float(dm_sec.get("당기사용원재료", 0.0) or 0.0)
    end = float(dm_sec.get("기말원재료", 0.0) or 0.0)

    rows.append(["직접재료비(참고)", "기초원재료", beg, ""])
    rows.append(["직접재료비(참고)", "당기원재료입고", inbound, ""])
    rows.append(["직접재료비", "당기사용원재료(원가반영)", used, ratio_str(safe_div(used, cogm)) if cogm else ""])
    rows.append(["직접재료비(참고)", "기말원재료", end, ""])
    rows.append(["직접재료비 소계(원가반영)", "직접재료비 소계", used, ratio_str(safe_div(used, cogm)) if cogm else ""])
    rows.append(["", "", 0.0, ""])

    dl_sum = sum_cost_section(cost_items, "직접노무비")
    for item, amt in cost_items.get("직접노무비", {}).items():
        amt = float(amt or 0)
        rows.append(["직접노무비", item, amt, ratio_str(safe_div(amt, cogm)) if cogm else ""])
    rows.append(["직접노무비 소계", "직접노무비 소계", dl_sum, ratio_str(safe_div(dl_sum, cogm)) if cogm else ""])
    rows.append(["", "", 0.0, ""])

    moh_sum = sum_cost_section(cost_items, "제조간접비")
    for item, amt in cost_items.get("제조간접비", {}).items():
        amt = float(amt or 0)
        rows.append(["제조간접비", item, amt, ratio_str(safe_div(amt, cogm)) if cogm else ""])
    rows.append(["제조간접비 소계", "제조간접비 소계", moh_sum, ratio_str(safe_div(moh_sum, cogm)) if cogm else ""])

    return pd.DataFrame(rows, columns=["구분", "항목", "금액", "당기제조원가 대비 비율"])


def df_pl_statement(p: PeriodInput, out: dict) -> pd.DataFrame:
    lines = [
        ("매출액", p.sales),
        ("매출원가", out["cogs"]),
        ("매출총이익", out["gross_profit"]),
        ("판관비(고정)", out["sga_fixed"]),
        ("판관비(변동)", out["sga_variable"]),
        ("판관비 합계", out["sga_total"]),
        ("영업이익", out["op_profit"]),
        ("영업외수익", out["nonop_income"]),
        ("영업외비용", out["nonop_expense"]),
        ("법인세차감전이익", out["pretax"]),
        ("법인세비용", out["tax"]),
        ("당기순이익", out["net_income"]),
    ]
    df = pd.DataFrame(lines, columns=["항목", "금액"])
    df["매출액 대비 비율"] = df["금액"].apply(lambda v: ratio_str(safe_div(float(v), p.sales)) if p.sales else "")
    return df


def df_bep(out: dict, sales_base: float) -> pd.DataFrame:
    # ✅ 이자 포함 BEP만
    interest_bep = float(out.get("interest_bep_sales", 0.0) or 0.0)
    return pd.DataFrame([
        ["변동비 합계(매출원가+판관비 변동)", out["variable_total"]],
        ["공헌이익", out["cm"]],
        ["공헌이익률", out["cm_ratio"]],
        ["고정비(판관비 고정)", out["sga_fixed"]],
        ["영업외비용(이자 등)", out["nonop_expense"]],
        ["이자 포함 BEP 매출액", interest_bep],
        ["이자 포함 BEP 대비 매출", sales_base - interest_bep],
        ["공헌이익률 등급", out["grade"]],
    ], columns=["항목", "값"])


def build_agg_pl_df(agg: dict) -> pd.DataFrame:
    sales_base = float(agg.get("sales", 0.0))
    df = pd.DataFrame([
        ["매출액", agg.get("sales", 0.0)],
        ["매출원가(합산)", agg.get("cogs", 0.0)],
        ["매출총이익", agg.get("gross_profit", 0.0)],
        ["판관비(고정)", agg.get("sga_fixed", 0.0)],
        ["판관비(변동)", agg.get("sga_variable", 0.0)],
        ["판관비 합계", agg.get("sga_total", 0.0)],
        ["영업이익", agg.get("op_profit", 0.0)],
        ["영업외수익", agg.get("nonop_income", 0.0)],
        ["영업외비용", agg.get("nonop_expense", 0.0)],
        ["법인세차감전이익", agg.get("pretax", 0.0)],
        ["법인세비용(월합산)", agg.get("tax", 0.0)],
        ["당기순이익", agg.get("net_income", 0.0)],
    ], columns=["항목", "금액"])
    df["매출액 대비 비율"] = df["금액"].apply(lambda v: ratio_str(safe_div(float(v), sales_base)) if sales_base else "")
    return df


# =============================
# Trend / Compare (회사/합산 대응)
# =============================
def aggregate_periods(payloads: list[dict]) -> dict:
    total = {
        "sales": 0.0,
        "cogs": 0.0,
        "gross_profit": 0.0,
        "sga_total": 0.0,
        "sga_fixed": 0.0,
        "sga_variable": 0.0,
        "op_profit": 0.0,
        "nonop_income": 0.0,
        "nonop_expense": 0.0,
        "pretax": 0.0,
        "tax": 0.0,
        "net_income": 0.0,
        "variable_total": 0.0,
        "months_included": [],
    }

    for payload in payloads:
        p = payload_to_period_input(payload)
        out = compute_all(p)

        total["sales"] += p.sales
        total["cogs"] += out["cogs"]
        total["gross_profit"] += out["gross_profit"]
        total["sga_total"] += out["sga_total"]
        total["sga_fixed"] += out["sga_fixed"]
        total["sga_variable"] += out["sga_variable"]
        total["op_profit"] += out["op_profit"]
        total["nonop_income"] += out["nonop_income"]
        total["nonop_expense"] += out["nonop_expense"]
        total["pretax"] += out["pretax"]
        total["tax"] += out["tax"]
        total["net_income"] += out["net_income"]
        total["variable_total"] += out["variable_total"]
        total["months_included"].append(int(payload.get("_month", 0)))

    cm = total["sales"] - total["variable_total"]
    cm_ratio = safe_div(cm, total["sales"])
    grade = grade_cm_ratio(cm_ratio)

    interest_bep_sales = ((total["sga_fixed"] + total["nonop_expense"]) / cm_ratio) if cm_ratio > 0 else 0.0

    total.update({
        "cm": cm,
        "cm_ratio": cm_ratio,
        "grade": grade,
        "interest_bep_sales": interest_bep_sales,
        "gap_sales_vs_interest_bep": total["sales"] - interest_bep_sales,
        "gross_profit_rate": safe_div(total["gross_profit"], total["sales"]),
        "op_margin": safe_div(total["op_profit"], total["sales"]),
    })
    return total


def aggregate_periods_group(year: int, months: list[int], companies: list[str]) -> tuple[dict, list[int]]:
    month_pack = load_periods_range_group(year, months, companies)

    payloads_merged_month = []
    included_months = []

    for pack in month_pack:
        m = int(pack.get("_month", 0))
        payloads = pack.get("_payloads", [])
        if not payloads:
            continue
        merged_payload = merge_payloads_for_group(payloads)
        merged_payload["_year"] = int(year)
        merged_payload["_month"] = int(m)
        payloads_merged_month.append(merged_payload)
        included_months.append(m)

    if not payloads_merged_month:
        return {}, []

    agg = aggregate_periods(payloads_merged_month)
    return agg, sorted(included_months)


# =============================
# 앱 시작 시 DB 스키마/정리 1회 보장
# =============================
_conn_boot = db_conn()
ensure_db_schema(_conn_boot)
_conn_boot.close()


# =============================
# UI (v0.6.3 형태 유지 + 회사 선택)
# =============================
st.title("제조원가·손익·BEP 경영분석 앱")

with st.sidebar:
    st.header("회사/연도/월 선택")

    company_selected = st.selectbox("회사", options=[GROUP_OPTION] + COMPANIES, index=0)

    now = datetime.now()
    year = st.number_input("연도", value=now.year, min_value=2000, max_value=2100, step=1)
    month = st.number_input("월", value=now.month, min_value=1, max_value=12, step=1)

    st.divider()

    saved = list_periods(company_selected if company_selected != GROUP_OPTION else None)
    if saved:
        st.caption("저장된 월")
        df_saved = pd.DataFrame(saved, columns=["company", "year", "month", "updated_at"])
        if company_selected != GROUP_OPTION:
            df_saved = df_saved[df_saved["company"] == company_selected]
        st.dataframe(df_saved, use_container_width=True, hide_index=True)

    st.divider()
    st.subheader("데이터 삭제")
    confirm = st.checkbox("정말 삭제합니다(되돌릴 수 없음).")

    if company_selected == GROUP_OPTION:
        if st.button("선택한 연/월 (전체회사) 삭제", use_container_width=True, disabled=not confirm):
            delete_period_all_companies(int(year), int(month))
            st.success("삭제 완료")
            st.rerun()
    else:
        if st.button("선택한 회사/연/월 삭제", use_container_width=True, disabled=not confirm):
            delete_period(company_selected, int(year), int(month))
            st.success("삭제 완료")
            st.rerun()

    st.divider()
    st.subheader("전체 초기화(모든 회사/모든 월)")
    if st.button("⚠️ DB 전체 데이터 완전 삭제", use_container_width=True, disabled=not confirm):
        purge_all_data()
        st.success("전체 데이터 삭제 완료")
        st.rerun()


# =============================
# defaults 로딩
# =============================
if company_selected == GROUP_OPTION:
    defaults = {}
else:
    defaults = load_period(company_selected, int(year), int(month))


def dget(key, fallback):
    return defaults.get(key, fallback)


sales_default = float(dget("sales", 0.0))
beg_fg_default = float(dget("beg_fg", 0.0))
end_fg_default = float(dget("end_fg", 0.0))

cost_items_default = dget("cost_items", None) or build_default_cost_items()
tax_mode_default = dget("tax_mode", "AUTO")
tax_manual_default = float(dget("tax_manual", 0.0))
include_local_default = bool(dget("include_local_tax", True))
auto_dm_default = bool(dget("auto_dm_used", True))

pkey = period_key(company_selected, int(year), int(month))
init_work_state_for_period(pkey, defaults)

tab1, tab2, tab3, tab4, tab5 = st.tabs(["입력", "결과", "누적 조회", "리포트 다운로드", "비교분석"])


# -----------------------------
# TAB1: INPUT
# -----------------------------
with tab1:
    st.subheader("1) 입력")

    if company_selected == GROUP_OPTION:
        st.info("‘전체(합산)’은 입력이 아니라 **회사별로 저장된 값을 자동 합산**해서 보여줍니다.\n\n좌측에서 회사 선택 후 입력/저장해 주세요.")
    else:
        left, right = st.columns(2)

        with left:
            st.markdown("### 매출 / 완제품 재고")
            sales = st.number_input("매출액", value=sales_default, step=1000.0, format="%.0f")
            beg_fg = st.number_input("기초 완제품 재고", value=beg_fg_default, step=1000.0, format="%.0f")
            end_fg = st.number_input("기말 완제품 재고", value=end_fg_default, step=1000.0, format="%.0f")

            st.markdown("### 제조원가명세서")
            st.info(
                "당기제조원가 = 직접재료비(당기사용원재료만) + 직접노무비 + 제조간접비\n\n"
                "직접재료비는 원재료만, 부재료/외주가공비/포장재/소모품은 제조간접비\n\n"
                "옵션 ON 시: 당기사용원재료 = (기초원재료 + 당기원재료입고 - 기말원재료)"
            )

            auto_dm_used = st.toggle("당기사용원재료 자동계산(기초+입고-기말, 음수면 0)", value=auto_dm_default)

            cost_items = {}
            for sec in ["직접재료비", "직접노무비", "제조간접비"]:
                with st.expander(sec, expanded=True if sec != "제조간접비" else False):
                    cost_items[sec] = {}

                    if sec == "직접재료비":
                        dm_sec_default = cost_items_default.get(sec, {})

                        beg = st.number_input(
                            "기초원재료", value=float(dm_sec_default.get("기초원재료", 0.0)),
                            step=1000.0, format="%.0f", help="참고용", key="dm_beg"
                        )

                        inbound_default = dm_sec_default.get("당기원재료입고", None)
                        if inbound_default is None:
                            inbound_default = dm_sec_default.get("당기원재료매입", 0.0)

                        inbound = st.number_input(
                            "당기원재료입고", value=float(inbound_default or 0.0),
                            step=1000.0, format="%.0f", help="입고 기준(실물/수불 기준)", key="dm_inbound"
                        )

                        end = st.number_input(
                            "기말원재료", value=float(dm_sec_default.get("기말원재료", 0.0)),
                            step=1000.0, format="%.0f", help="참고용", key="dm_end"
                        )

                        raw_used = float(beg) + float(inbound) - float(end)
                        auto_used = max(0.0, raw_used)

                        used_default = float(dm_sec_default.get("당기사용원재료", 0.0))
                        used_val = auto_used if auto_dm_used else st.number_input(
                            "당기사용원재료(원가반영)", value=used_default,
                            step=1000.0, format="%.0f",
                            help="원가에 반영되는 직접재료비(원재료 사용액)", disabled=auto_dm_used,
                            key="dm_used_manual"
                        )

                        if auto_dm_used:
                            if raw_used < 0:
                                st.warning(f"자동계산(기초+입고-기말) = {to_money(raw_used)} → 음수이므로 0으로 보정되었습니다.")
                            st.caption(f"자동계산된 당기사용원재료(원가반영): {to_money(used_val)}")

                        cost_items[sec]["기초원재료"] = float(beg)
                        cost_items[sec]["당기원재료입고"] = float(inbound)
                        cost_items[sec]["기말원재료"] = float(end)
                        cost_items[sec]["당기사용원재료"] = float(used_val)

                    else:
                        item_names = list(cost_items_default.get(sec, {}).keys())
                        for i, item in enumerate(item_names):
                            val = float(cost_items_default.get(sec, {}).get(item, 0.0))
                            amt = st.number_input(item, value=val, step=1000.0, format="%.0f", key=f"cost_{sec}_{i}")
                            cost_items[sec][item] = float(amt)

                        st.caption("필요 시 항목 추가(선택)")
                        new_name = st.text_input(f"{sec} 추가 항목명", value="", key=f"new_{sec}")
                        if new_name.strip():
                            if new_name not in cost_items[sec]:
                                amt = st.number_input(
                                    f"[추가] {new_name}", value=0.0, step=1000.0, format="%.0f",
                                    key=f"cost_{sec}_add_{abs(hash(new_name))}"
                                )
                                cost_items[sec][new_name] = float(amt)

        with right:
            st.markdown("### 판매관리비 (고정비/변동비)")

            sga_rows_work = st.session_state.get("sga_rows_work", build_default_sga_rows())
            sga_rows = []

            for i, r in enumerate(sga_rows_work):
                c1, c2, c3 = st.columns([3, 1.5, 2.5])
                with c1:
                    item = st.text_input("항목", value=r.get("item", ""), key=f"sga_item_{i}")
                with c2:
                    typ0 = normalize_sga_type(r.get("type", "고정"))
                    typ = st.selectbox("구분", ["고정", "변동"], index=["고정", "변동"].index(typ0), key=f"sga_type_{i}")
                with c3:
                    amt = st.number_input("금액", value=float(r.get("amount", 0.0)), step=1000.0, format="%.0f", key=f"sga_amt_{i}")
                sga_rows.append({"item": item, "type": typ, "amount": float(amt)})

            st.session_state["sga_rows_work"] = sga_rows

            st.divider()
            st.caption("판관비 항목 추가(선택)")
            add_name = st.text_input("새 판관비 항목명", value="", key="sga_new_name")
            add_type = st.selectbox("새 항목 구분", ["고정", "변동"], key="sga_new_type")
            add_amt = st.number_input("새 항목 금액", value=0.0, step=1000.0, format="%.0f", key="sga_new_amt")

            if st.button("판관비 항목 추가", use_container_width=True):
                if add_name.strip():
                    st.session_state["sga_rows_work"] = st.session_state["sga_rows_work"] + [{
                        "item": add_name.strip(),
                        "type": add_type,
                        "amount": float(add_amt)
                    }]
                    st.success("추가됨! (저장 버튼을 누르면 다음에도 유지됩니다)")
                    st.rerun()
                else:
                    st.warning("항목명을 입력해 주세요.")

            st.markdown("### 영업외 수익/비용")
            with st.expander("영업외수익", expanded=False):
                nonop_income_rows = []
                for i, r in enumerate(st.session_state.get("nonop_income_work", build_default_nonop_rows(DEFAULT_NONOP_INCOME))):
                    c1, c2 = st.columns([3, 2])
                    with c1:
                        item = st.text_input("항목", value=r.get("item", ""), key=f"noi_item_{i}")
                    with c2:
                        amt = st.number_input("금액", value=float(r.get("amount", 0.0)), step=1000.0, format="%.0f", key=f"noi_amt_{i}")
                    nonop_income_rows.append({"item": item, "amount": float(amt)})
                st.session_state["nonop_income_work"] = nonop_income_rows

            with st.expander("영업외비용", expanded=False):
                nonop_expense_rows = []
                for i, r in enumerate(st.session_state.get("nonop_expense_work", build_default_nonop_rows(DEFAULT_NONOP_EXPENSE))):
                    c1, c2 = st.columns([3, 2])
                    with c1:
                        item = st.text_input("항목", value=r.get("item", ""), key=f"noe_item_{i}")
                    with c2:
                        amt = st.number_input("금액", value=float(r.get("amount", 0.0)), step=1000.0, format="%.0f", key=f"noe_amt_{i}")
                    nonop_expense_rows.append({"item": item, "amount": float(amt)})
                st.session_state["nonop_expense_work"] = nonop_expense_rows

            st.markdown("### 법인세 설정")
            tax_mode = st.selectbox("법인세 반영 방식", ["AUTO", "MANUAL", "OFF"],
                                    index=["AUTO", "MANUAL", "OFF"].index(tax_mode_default))
            include_local_tax = st.checkbox("지방소득세(국세의 10%) 포함(추정)", value=include_local_default)
            tax_manual = 0.0
            if tax_mode == "MANUAL":
                tax_manual = st.number_input("법인세비용(수동 입력)", value=tax_manual_default, step=1000.0, format="%.0f")

        payload = {
            "sales": float(sales),
            "beg_fg": float(beg_fg),
            "end_fg": float(end_fg),
            "cost_items": cost_items,
            "sga_rows": st.session_state.get("sga_rows_work", sga_rows),
            "nonop_income_rows": st.session_state.get("nonop_income_work", nonop_income_rows),
            "nonop_expense_rows": st.session_state.get("nonop_expense_work", nonop_expense_rows),
            "tax_mode": tax_mode,
            "tax_manual": float(tax_manual) if tax_mode == "MANUAL" else float(tax_manual_default),
            "include_local_tax": bool(include_local_tax),
            "auto_dm_used": bool(auto_dm_used),
        }

        s1, s2 = st.columns(2)
        with s1:
            if st.button("저장", use_container_width=True):
                save_period(company_selected, int(year), int(month), payload)
                st.success(f"저장 완료 ({company_selected})")
                st.session_state["_work_period_key"] = period_key(company_selected, int(year), int(month))
                st.session_state["sga_rows_work"] = payload["sga_rows"]
                st.session_state["nonop_income_work"] = payload["nonop_income_rows"]
                st.session_state["nonop_expense_work"] = payload["nonop_expense_rows"]
        with s2:
            if st.button("초기화(이번 달)", use_container_width=True):
                save_period(company_selected, int(year), int(month), {})
                st.warning("초기화 완료")
                st.session_state["_work_period_key"] = None
                st.rerun()


# -----------------------------
# Current compute (선택월)
# -----------------------------
if company_selected == GROUP_OPTION:
    packs = load_periods_range_group(int(year), [int(month)], COMPANIES)
    pack = packs[0] if packs else {"_payloads": []}
    payloads = pack.get("_payloads", [])

    if payloads:
        p, out = compute_all_group(payloads)
    else:
        p = PeriodInput(
            0, 0, 0,
            build_default_cost_items(),
            build_default_sga_rows(),
            build_default_nonop_rows(DEFAULT_NONOP_INCOME),
            build_default_nonop_rows(DEFAULT_NONOP_EXPENSE),
            "OFF", 0.0, True
        )
        out = compute_all(p)
else:
    curr_payload = load_period(company_selected, int(year), int(month)) or {}
    p = payload_to_period_input(curr_payload)
    out = compute_all(p)

df_cost = df_cost_statement(p.cost_items, out["cogm"])
df_pl = df_pl_statement(p, out)
df_bep_tbl = df_bep(out, p.sales)


# -----------------------------
# TAB2: RESULT
# -----------------------------
with tab2:
    st.subheader("2) 결과 (선택월)")

    if company_selected == GROUP_OPTION:
        st.caption(f"합산 기준 회사: {', '.join(COMPANIES)}")

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("당기제조원가", to_money(out["cogm"]))
    c2.metric("매출원가", to_money(out["cogs"]))
    c3.metric("영업이익", to_money(out["op_profit"]))
    c4.metric("당기순이익", to_money(out["net_income"]))

    st.divider()

    c5, c6, c7, c8 = st.columns(4)
    c5.metric("공헌이익률", f"{out['cm_ratio']*100:,.1f}%")
    c6.metric("공헌이익률 판정", out["grade"])
    c7.metric("이자 포함 BEP 매출액", to_money(out.get("interest_bep_sales", 0.0)))
    c8.metric("이자 포함 BEP 대비 매출", to_money(out.get("gap_sales_vs_interest_bep", 0.0)))

    left, right = st.columns(2)
    with left:
        st.markdown("### 손익계산서 요약")
        st.dataframe(df_pl, use_container_width=True, hide_index=True, height=520)

        st.markdown("### 이자 포함 BEP / 공헌이익")
        st.dataframe(df_bep_tbl, use_container_width=True, hide_index=True, height=320)

    with right:
        st.markdown("### 제조원가명세서 상세")
        st.dataframe(df_cost, use_container_width=True, hide_index=True, height=820)


# -----------------------------
# TAB3: CUMULATIVE
# -----------------------------
with tab3:
    st.subheader("3) 누적 조회 (저장된 월만 합산)")

    view_type = st.selectbox("조회 기준", ["월별(선택월)", "분기누적(QTD)", "연도누적(YTD)", "연도전체(저장된 월 합산)"])
    selected_year = int(year)
    selected_month = int(month)

    months = months_for_period(selected_month, view_type)
    if view_type == "월별(선택월)":
        period_label = f"{selected_year}-{selected_month:02d}"
    elif view_type == "분기누적(QTD)":
        q = (selected_month - 1) // 3 + 1
        period_label = f"{selected_year} Q{q} (저장분)"
    elif view_type == "연도누적(YTD)":
        period_label = f"{selected_year} YTD (저장분)"
    else:
        period_label = f"{selected_year} 연도전체(저장분)"

    if company_selected == GROUP_OPTION:
        agg, included = aggregate_periods_group(selected_year, months, COMPANIES)
        if not agg:
            st.warning("선택한 기간에 저장된 데이터가 없습니다. (회사별 입력 탭에서 월별 저장 후 다시 확인)")
        else:
            expected = set(months)
            missing = sorted(list(expected - set(included)))
            if missing:
                st.info(f"저장되지 않은 월은 누적에서 제외됩니다: {', '.join(map(str, missing))}")

            a1, a2, a3, a4 = st.columns(4)
            a1.metric("누적 매출액", to_money(agg["sales"]))
            a2.metric("누적 영업이익", to_money(agg["op_profit"]))
            a3.metric("누적 당기순이익", to_money(agg["net_income"]))
            a4.metric("누적 공헌이익률", f"{agg['cm_ratio']*100:,.1f}% ({agg['grade']})")

            b1, b2 = st.columns(2)
            b1.metric("누적 이자 포함 BEP 매출액", to_money(agg["interest_bep_sales"]))
            b2.metric("누적 이자 포함 BEP 대비 매출", to_money(agg["gap_sales_vs_interest_bep"]))

            st.caption(f"기간: {period_label} | 포함 월: {', '.join([str(m) for m in included])}")

            df_agg_pl = build_agg_pl_df(agg)
            st.markdown("### 누적 손익계산서 요약")
            st.dataframe(df_agg_pl, use_container_width=True, hide_index=True, height=420)

            st.markdown("### 누적 이자 포함 BEP / 공헌이익")
            df_agg_bep = pd.DataFrame([
                ["변동비 합계(매출원가+판관비 변동)", agg["variable_total"]],
                ["공헌이익", agg["cm"]],
                ["공헌이익률", agg["cm_ratio"]],
                ["고정비(판관비 고정)", agg["sga_fixed"]],
                ["영업외비용(이자 등)", agg["nonop_expense"]],
                ["이자 포함 BEP 매출액", agg["interest_bep_sales"]],
                ["이자 포함 BEP 대비 매출", agg["gap_sales_vs_interest_bep"]],
                ["공헌이익률 등급", agg["grade"]],
            ], columns=["항목", "값"])
            st.dataframe(df_agg_bep, use_container_width=True, hide_index=True, height=320)

    else:
        payloads = load_periods_range(company_selected, selected_year, months)
        if not payloads:
            st.warning("선택한 기간에 저장된 데이터가 없습니다. (입력 탭에서 저장 후 다시 확인)")
        else:
            expected = set(months)
            included_m = set(int(pp.get("_month", 0)) for pp in payloads)
            missing = sorted(list(expected - included_m))
            if missing:
                st.info(f"저장되지 않은 월은 누적에서 제외됩니다: {', '.join(map(str, missing))}")

            agg = aggregate_periods(payloads)

            a1, a2, a3, a4 = st.columns(4)
            a1.metric("누적 매출액", to_money(agg["sales"]))
            a2.metric("누적 영업이익", to_money(agg["op_profit"]))
            a3.metric("누적 당기순이익", to_money(agg["net_income"]))
            a4.metric("누적 공헌이익률", f"{agg['cm_ratio']*100:,.1f}% ({agg['grade']})")

            b1, b2 = st.columns(2)
            b1.metric("누적 이자 포함 BEP 매출액", to_money(agg["interest_bep_sales"]))
            b2.metric("누적 이자 포함 BEP 대비 매출", to_money(agg["gap_sales_vs_interest_bep"]))

            st.caption(f"기간: {period_label} | 포함 월: {', '.join([str(m) for m in sorted([x for x in agg['months_included'] if x])])}")

            df_agg_pl = build_agg_pl_df(agg)
            st.markdown("### 누적 손익계산서 요약")
            st.dataframe(df_agg_pl, use_container_width=True, hide_index=True, height=420)

            st.markdown("### 누적 이자 포함 BEP / 공헌이익")
            df_agg_bep = pd.DataFrame([
                ["변동비 합계(매출원가+판관비 변동)", agg["variable_total"]],
                ["공헌이익", agg["cm"]],
                ["공헌이익률", agg["cm_ratio"]],
                ["고정비(판관비 고정)", agg["sga_fixed"]],
                ["영업외비용(이자 등)", agg["nonop_expense"]],
                ["이자 포함 BEP 매출액", agg["interest_bep_sales"]],
                ["이자 포함 BEP 대비 매출", agg["gap_sales_vs_interest_bep"]],
                ["공헌이익률 등급", agg["grade"]],
            ], columns=["항목", "값"])
            st.dataframe(df_agg_bep, use_container_width=True, hide_index=True, height=320)


# -----------------------------
# TAB4: 리포트 다운로드 (UI 자리만 유지)
# -----------------------------

def build_report_html(title: str, period_label: str, sales_base: float, out: dict,
                      df_pl: pd.DataFrame, df_cost: pd.DataFrame | None,
                      tax_mode: str, tax_info: dict | None) -> str:
    tax_note = ""
    if tax_info:
        d = tax_info.get("detail", {})
        tax_note = f"""
        <p><b>법인세(표시/추정)</b><br/>
        연환산 과세표준: {tax_info.get("annualized_tax_base", 0):,.0f}원<br/>
        국세 법인세: {d.get("national_cit", 0):,.0f}원 / 지방소득세(가정): {d.get("local_income_tax", 0):,.0f}원<br/>
        ({d.get("assumption", "")})
        </p>
        """
    else:
        if tax_mode == "OFF":
            tax_note = "<p><b>법인세:</b> 미반영(0원)</p>"
        elif tax_mode == "MANUAL":
            tax_note = "<p><b>법인세:</b> 수동 입력값 반영</p>"

    cost_block = ""
    if df_cost is not None:
        cost_block = f"""
        <h3>제조원가명세서(당기제조원가 대비 비율)</h3>
        <p><small>※ 직접재료비는 '당기사용원재료'만 원가에 반영됩니다. 기초/입고/기말은 참고 표시입니다.</small></p>
        {df_cost.to_html(index=False)}
        """

    html = f"""
    <h2>{title}</h2>
    <p><b>기간:</b> {period_label}</p>

    <h3>핵심 KPI</h3>
    <ul>
      <li>매출액: {sales_base:,.0f}</li>
      <li>영업이익: {out.get('op_profit',0):,.0f} (영업이익률 {out.get('op_margin',0)*100:,.1f}%)</li>
      <li>공헌이익률: {out.get('cm_ratio',0)*100:,.1f}% → <b>{out.get('grade','')}</b></li>
      <li>이자포함 BEP 매출액: {out.get('interest_bep_sales',0):,.0f}</li>
      <li>이자포함 BEP 대비 매출: {out.get('gap_sales_vs_interest_bep',0):,.0f}</li>
    </ul>

    {tax_note}

    <h3>손익계산서 요약</h3>
    {df_pl.to_html(index=False)}

    {cost_block}

    <hr/>
    <p><small>
    ※ 본 리포트는 경영관리 목적의 자동 산출 결과입니다.
    특히 법인세는 추정/합산 방식에 따라 실제 신고세액과 차이가 있을 수 있습니다.
    </small></p>
    """
    return html


def make_excel_report(df_pl: pd.DataFrame, df_cost: pd.DataFrame | None,
                      df_bep_tbl: pd.DataFrame, extra_sheets: dict[str, pd.DataFrame] | None = None) -> bytes:
    """
    ✅ 엑셀 리포트 생성
    - Streamlit Cloud 환경에서 openpyxl 미설치로 오류가 나는 경우가 많아서
      openpyxl → xlsxwriter 순서로 자동 폴백합니다.
    - 둘 다 없으면 RuntimeError 를 발생시키며, UI 쪽에서 try/except로 안내합니다.
    """
    buf = io.BytesIO()

    engine = None
    try:
        import openpyxl  # noqa: F401
        engine = "openpyxl"
    except Exception:
        try:
            import xlsxwriter  # noqa: F401
            engine = "xlsxwriter"
        except Exception:
            engine = None

    if engine is None:
        raise RuntimeError(
            "엑셀 저장 엔진이 없습니다. requirements.txt에 'openpyxl' 또는 'xlsxwriter'를 추가해 주세요."
        )

    with pd.ExcelWriter(buf, engine=engine) as writer:
        df_pl.to_excel(writer, index=False, sheet_name="손익요약")
        df_bep_tbl.to_excel(writer, index=False, sheet_name="이자포함BEP")
        if df_cost is not None:
            df_cost.to_excel(writer, index=False, sheet_name="제조원가")
        if extra_sheets:
            for name, df in extra_sheets.items():
                df.to_excel(writer, index=False, sheet_name=name[:31])

    buf.seek(0)
    return buf.getvalue()



def make_pdf_report(title: str, period_label: str, sales_base: float, out: dict,
                    df_pl: pd.DataFrame, tax_mode: str, tax_info: dict | None) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4

    font = register_korean_font_if_possible()

    c.setFont(font, 14)
    c.drawString(18 * mm, height - 18 * mm, title)

    c.setFont(font, 11)
    y = height - 28 * mm
    line = 6.5 * mm

    c.drawString(18 * mm, y, f"기간: {period_label}")
    y -= line

    c.drawString(18 * mm, y, f"매출액: {to_money(sales_base)}")
    y -= line
    c.drawString(18 * mm, y, f"영업이익: {to_money(out.get('op_profit',0))}")
    y -= line
    c.drawString(18 * mm, y, f"공헌이익률: {out.get('cm_ratio',0)*100:,.1f}% ({out.get('grade','')})")
    y -= line
    c.drawString(18 * mm, y, f"이자포함 BEP 매출액: {to_money(out.get('interest_bep_sales',0))}")
    y -= (line * 1.2)

    if tax_info:
        d = tax_info.get("detail", {})
        c.drawString(18 * mm, y, f"법인세(표시): {tax_info.get('note','')}")
        y -= line
        c.drawString(18 * mm, y, f"- 합계 {to_money(d.get('total_tax',0))}")
        y -= (line * 1.2)
    else:
        if tax_mode == "OFF":
            c.drawString(18 * mm, y, "법인세: 미반영(0원)")
            y -= (line * 1.2)
        elif tax_mode == "MANUAL":
            c.drawString(18 * mm, y, "법인세: 수동 입력값 반영")
            y -= (line * 1.2)

    c.setFont(font, 12)
    c.drawString(18 * mm, y, "손익계산서 요약")
    y -= (line * 1.2)

    c.setFont(font, 10)
    for _, row in df_pl.iterrows():
        item = str(row["항목"])
        amt = float(row["금액"])
        ratio = str(row.get("매출액 대비 비율", ""))

        c.drawString(20 * mm, y, item)
        c.drawRightString(width - 50 * mm, y, f"{amt:,.0f}")
        c.drawRightString(width - 18 * mm, y, ratio)

        y -= line
        if y < 18 * mm:
            c.showPage()
            c.setFont(font, 10)
            y = height - 18 * mm

    c.showPage()
    c.save()
    buf.seek(0)
    return buf.getvalue()


# -----------------------------
# TAB4: REPORT DOWNLOAD (PDF / Excel / HTML)
# -----------------------------
with tab4:
    st.subheader("4) 리포트 다운로드")

    # ✅ 현재 화면(선택월) 기준 리포트 생성
    report_title = f"경영분석 리포트 ({company_selected} {int(year)}-{int(month):02d})"
    period_label = f"{int(year)}-{int(month):02d}"
    sales_base = float(p.sales)

    html = build_report_html(
        title=report_title,
        period_label=period_label,
        sales_base=sales_base,
        out=out,
        df_pl=df_pl,
        df_cost=df_cost,
        tax_mode=getattr(p, "tax_mode", "OFF"),
        tax_info=out.get("tax_info", None),
    )

    # ✅ 엑셀: 손익요약 + BEP + 제조원가
    try:
        excel_bytes = make_excel_report(
            df_pl=df_pl,
            df_cost=df_cost,
            df_bep_tbl=df_bep_tbl,
            extra_sheets=None,
        )
        excel_ok = True
        excel_err = ""
    except Exception as e:
        excel_bytes = b""
        excel_ok = False
        excel_err = str(e)

    # ✅ PDF: 손익요약 중심(간단)
    pdf_bytes = make_pdf_report(
        title=report_title,
        period_label=period_label,
        sales_base=sales_base,
        out=out,
        df_pl=df_pl,
        tax_mode=getattr(p, "tax_mode", "OFF"),
        tax_info=out.get("tax_info", None),
    )

    c1, c2, c3 = st.columns(3)

    with c1:
        if not excel_ok:
            st.warning(f"엑셀 다운로드를 만들 수 없습니다: {excel_err}")
            st.caption("Streamlit Cloud라면 requirements.txt에 openpyxl 또는 xlsxwriter를 추가해 주세요.")
        st.download_button(
            "📥 엑셀 다운로드 (.xlsx)",
            data=excel_bytes,
            file_name=f"report_{company_selected}_{int(year)}_{int(month):02d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            disabled=not excel_ok,
        )

    with c2:
        st.download_button(
            "📥 PDF 다운로드 (.pdf)",
            data=pdf_bytes,
            file_name=f"report_{company_selected}_{int(year)}_{int(month):02d}.pdf",
            mime="application/pdf",
            use_container_width=True,
        )

    with c3:
        st.download_button(
            "📥 HTML 다운로드 (.html)",
            data=html.encode("utf-8"),
            file_name=f"report_{company_selected}_{int(year)}_{int(month):02d}.html",
            mime="text/html",
            use_container_width=True,
        )

    st.caption("※ 리포트는 현재 화면(선택월) 계산 결과를 그대로 내보냅니다.")


# -----------------------------
# TAB5: 비교분석
# -----------------------------
with tab5:
    st.subheader("회사별 비교분석")

    months = [int(month)]
    rows = []

    packs = load_periods_range_group(int(year), months, COMPANIES)

    for pack in packs:
        for company, payload in pack.get("_payloads", []):
            p1 = payload_to_period_input(payload)
            out1 = compute_all(p1)
            rows.append({
                "회사": company,
                "매출액": p1.sales,
                "영업이익": out1["op_profit"],
                "당기순이익": out1["net_income"],
                "공헌이익률": out1["cm_ratio"],
                "고정비": out1["sga_fixed"],
                "영업외비용": out1["nonop_expense"],
                "이자포함BEP매출액": out1.get("interest_bep_sales", 0.0),
                "이자포함BEP대비매출": out1.get("gap_sales_vs_interest_bep", 0.0),
            })

    if not rows:
        st.warning("비교할 데이터가 없습니다.")
    else:
        df = pd.DataFrame(rows)

        st.markdown("### 회사별 KPI 비교")
        st.dataframe(df, use_container_width=True, hide_index=True)

        st.markdown("### 공헌이익률 랭킹")
        st.dataframe(
            df.sort_values("공헌이익률", ascending=False)[["회사", "공헌이익률"]],
            use_container_width=True, hide_index=True
        )

        st.markdown("### 이자 포함 BEP 미달 회사 (고정비+영업외비용 기준)")
        st.dataframe(
            df[df["이자포함BEP대비매출"] < 0][["회사", "이자포함BEP대비매출"]],
            use_container_width=True, hide_index=True
        )


st.divider()
st.caption(f"버전: v{APP_VERSION} | 개발자: {DEVELOPER_NAME} | {COPYRIGHT_TEXT}")
